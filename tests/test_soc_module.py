"""
Comprehensive Test Suite for SOC Processing Module
Apollo Revenue Audit - SOC Data Processing Module
"""

import os
import shutil
import unittest
import openpyxl
import fitz  # PyMuPDF
import sqlite3
import json

from soc_module.models import SOCRecord
from soc_module.config import MappingConfigManager
from soc_module.excel_parser import ExcelSOCParser
from soc_module.pdf_parser import PDFSOCParser
from soc_module.validator import SOCValidator
from soc_module.tariff_integrator import TariffIntegrator
from soc_module.soc_manager import SOCProcessingManager

TEST_OUTPUT_DIR = "tests/test_artifacts"

class TestSOCModule(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        os.makedirs(TEST_OUTPUT_DIR, exist_ok=True)
        cls.test_db_path = os.path.join(TEST_OUTPUT_DIR, "test_revenue_audit.db")
        if os.path.exists(cls.test_db_path):
            os.remove(cls.test_db_path)
        cls.manager = SOCProcessingManager(db_path=cls.test_db_path)

    @classmethod
    def tearDownClass(cls):
        if os.path.exists(TEST_OUTPUT_DIR):
            shutil.rmtree(TEST_OUTPUT_DIR, ignore_errors=True)

    # -------------------------------------------------------------
    # TEST 1: Standard Excel SOC
    # -------------------------------------------------------------
    def test_01_standard_excel_soc(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_01_standard.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TariffMaster"
        ws.append(["SERVICEID", "SERVICENAME", "DEPTNAME", "SERVICETYPENAME", "FINALTARIFF", "ALIASCODE"])
        ws.append(["101", "CONSULTATION OPD", "General Medicine", "Consultation", 500, "CON01"])
        ws.append(["102", "X-RAY CHEST PA", "Radiology", "Investigation", 750, "RAD02"])
        ws.append(["103", "CBC WITH ESR", "Pathology", "Investigation", 450, "LAB03"])
        wb.save(file_path)

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["validation"]["valid_count"], 3)
        self.assertEqual(res["validation"]["invalid_count"], 0)
        self.assertEqual(res["records"][0]["service_code"], "101")
        self.assertEqual(res["records"][0]["rate"], 500.0)

    # -------------------------------------------------------------
    # TEST 2: Excel with Blank Rows & Header Banners
    # -------------------------------------------------------------
    def test_02_excel_with_blank_rows(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_02_blank_rows.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Apollo Hospitals Enterprise SOC Schedule", None, None])
        ws.append([None, None, None])  # Blank
        ws.append(["Effective Date: 2026-04-01", None, None])
        ws.append(["Service ID", "Service Description", "Amount"])
        ws.append(["201", "MRI BRAIN PLAIN", 6500])
        ws.append([None, None, None])  # Blank row in data
        ws.append(["202", "CT SCAN HEAD", 3500])
        ws.append(["", " ", "  "])    # Whitespace row
        ws.append(["203", "ULTRASOUND ABDOMEN", 1800])
        wb.save(file_path)

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["validation"]["valid_count"], 3)
        codes = [r["service_code"] for r in res["records"]]
        self.assertListEqual(codes, ["201", "202", "203"])

    # -------------------------------------------------------------
    # TEST 3: Excel with Different / Custom Column Names
    # -------------------------------------------------------------
    def test_03_excel_with_different_column_names(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_03_custom_cols.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Item Code", "Item Description", "Specialty Group", "Classification", "Negotiated Rate", "UOM"])
        ws.append(["C-991", "ANGIOGRAPHY CORONARY", "Cardiology", "Procedure", 18500, "Per Case"])
        ws.append(["C-992", "ANGIOPLASTY SINGLE VESSEL", "Cardiology", "Procedure", 95000, "Per Case"])
        wb.save(file_path)

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["validation"]["valid_count"], 2)
        self.assertEqual(res["records"][0]["service_code"], "C-991")
        self.assertEqual(res["records"][0]["department"], "Cardiology")
        self.assertEqual(res["records"][0]["rate"], 18500.0)
        self.assertEqual(res["records"][0]["unit"], "Per Case")

    # -------------------------------------------------------------
    # TEST 4: PDF with a Standard Tariff Table
    # -------------------------------------------------------------
    def test_04_pdf_standard_table(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_04_standard_table.pdf")
        doc = fitz.open()
        page = doc.new_page()
        # Draw explicit gridlines for pdfplumber table extraction
        # Title
        page.insert_text((50, 40), "APOLLO HOSPITALS GUWAHATI - TARIFF SCHEDULE", fontsize=12)
        
        # Table lines
        y_positions = [60, 85, 110, 135, 160]
        x_positions = [50, 150, 400, 500]
        
        for y in y_positions:
            page.draw_line(fitz.Point(50, y), fitz.Point(500, y))
        for x in x_positions:
            page.draw_line(fitz.Point(x, 60), fitz.Point(x, 160))

        # Text in cells
        page.insert_text((55, 78), "SERVICE CODE", fontsize=9)
        page.insert_text((155, 78), "DESCRIPTION", fontsize=9)
        page.insert_text((405, 78), "TARIFF RATE", fontsize=9)

        page.insert_text((55, 103), "3001", fontsize=9)
        page.insert_text((155, 103), "BLOOD SUGAR FASTING", fontsize=9)
        page.insert_text((405, 103), "120.00", fontsize=9)

        page.insert_text((55, 128), "3002", fontsize=9)
        page.insert_text((155, 128), "SERUM CREATININE", fontsize=9)
        page.insert_text((405, 128), "220.00", fontsize=9)

        page.insert_text((55, 153), "3003", fontsize=9)
        page.insert_text((155, 153), "LIVER FUNCTION TEST", fontsize=9)
        page.insert_text((405, 153), "850.00", fontsize=9)

        doc.save(file_path)
        doc.close()

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "success")
        self.assertGreaterEqual(res["validation"]["valid_count"], 3)
        self.assertEqual(res["records"][0]["service_code"], "3001")
        self.assertEqual(res["records"][0]["rate"], 120.0)

    # -------------------------------------------------------------
    # TEST 5: PDF with Multiple Pages
    # -------------------------------------------------------------
    def test_05_pdf_multiple_pages(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_05_multi_page.pdf")
        doc = fitz.open()

        # Page 1
        p1 = doc.new_page()
        for y in [50, 75, 100]: p1.draw_line(fitz.Point(50, y), fitz.Point(450, y))
        for x in [50, 150, 350, 450]: p1.draw_line(fitz.Point(x, 50), fitz.Point(x, 100))
        p1.insert_text((55, 68), "SERVICE_ID", fontsize=9)
        p1.insert_text((155, 68), "SERVICE_NAME", fontsize=9)
        p1.insert_text((355, 68), "RATE", fontsize=9)
        p1.insert_text((55, 93), "P1-01", fontsize=9)
        p1.insert_text((155, 93), "PAGE 1 PROCEDURE A", fontsize=9)
        p1.insert_text((355, 93), "1500", fontsize=9)

        # Page 2
        p2 = doc.new_page()
        for y in [50, 75, 100]: p2.draw_line(fitz.Point(50, y), fitz.Point(450, y))
        for x in [50, 150, 350, 450]: p2.draw_line(fitz.Point(x, 50), fitz.Point(x, 100))
        p2.insert_text((55, 68), "SERVICE_ID", fontsize=9)
        p2.insert_text((155, 68), "SERVICE_NAME", fontsize=9)
        p2.insert_text((355, 68), "RATE", fontsize=9)
        p2.insert_text((55, 93), "P2-01", fontsize=9)
        p2.insert_text((155, 93), "PAGE 2 PROCEDURE B", fontsize=9)
        p2.insert_text((355, 93), "2500", fontsize=9)

        doc.save(file_path)
        doc.close()

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "success")
        self.assertGreaterEqual(res["validation"]["valid_count"], 2)
        codes = [r["service_code"] for r in res["records"]]
        self.assertIn("P1-01", codes)
        self.assertIn("P2-01", codes)

    # -------------------------------------------------------------
    # TEST 6: PDF with Multiple Tables
    # -------------------------------------------------------------
    def test_06_pdf_multiple_tables(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_06_multi_table.pdf")
        doc = fitz.open()
        p = doc.new_page()

        # Table 1: Consultations
        for y in [50, 75, 100]: p.draw_line(fitz.Point(50, y), fitz.Point(450, y))
        for x in [50, 150, 350, 450]: p.draw_line(fitz.Point(x, 50), fitz.Point(x, 100))
        p.insert_text((55, 68), "CODE", fontsize=9)
        p.insert_text((155, 68), "CONSULTATION", fontsize=9)
        p.insert_text((355, 68), "FEES", fontsize=9)
        p.insert_text((55, 93), "T1-CON", fontsize=9)
        p.insert_text((155, 93), "CARDIOLOGY OPD", fontsize=9)
        p.insert_text((355, 93), "800", fontsize=9)

        # Table 2: Investigations
        for y in [150, 175, 200]: p.draw_line(fitz.Point(50, y), fitz.Point(450, y))
        for x in [50, 150, 350, 450]: p.draw_line(fitz.Point(x, 150), fitz.Point(x, 200))
        p.insert_text((55, 168), "CODE", fontsize=9)
        p.insert_text((155, 168), "INVESTIGATION", fontsize=9)
        p.insert_text((355, 168), "FEES", fontsize=9)
        p.insert_text((55, 193), "T2-ECG", fontsize=9)
        p.insert_text((155, 193), "12 LEAD ECG", fontsize=9)
        p.insert_text((355, 193), "350", fontsize=9)

        doc.save(file_path)
        doc.close()

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "success")
        self.assertGreaterEqual(res["validation"]["valid_count"], 2)
        codes = [r["service_code"] for r in res["records"]]
        self.assertIn("T1-CON", codes)
        self.assertIn("T2-ECG", codes)

    # -------------------------------------------------------------
    # TEST 7: Missing Mandatory Fields
    # -------------------------------------------------------------
    def test_07_missing_mandatory_fields(self):
        records = [
            {"service_code": "", "description": "No Code Item", "rate": 500},
            {"service_code": "1002", "description": "", "rate": 600},
            {"service_code": "1003", "description": "Valid Item", "rate": 700}
        ]
        validator = SOCValidator()
        report = validator.validate_batch(records)
        self.assertEqual(report["total_records"], 3)
        self.assertEqual(report["valid_count"], 1)
        self.assertEqual(report["invalid_count"], 2)

    # -------------------------------------------------------------
    # TEST 8: Duplicate Tariff Records
    # -------------------------------------------------------------
    def test_08_duplicate_tariff_records(self):
        records = [
            {"service_code": "DUP001", "description": "First Definition", "rate": 500},
            {"service_code": "DUP001", "description": "Second Revised Definition", "rate": 550},
            {"service_code": "DUP002", "description": "Unique Item", "rate": 1000}
        ]
        validator = SOCValidator()
        report = validator.validate_batch(records)
        self.assertEqual(report["valid_count"], 3) # Both valid, but warning logged
        self.assertEqual(report["warning_count"], 1)
        self.assertIn("Duplicate service code 'DUP001'", report["warnings"][0]["warnings"][0])

    # -------------------------------------------------------------
    # TEST 9: Invalid Rates (Negative, Non-Numeric, NaN)
    # -------------------------------------------------------------
    def test_09_invalid_rates(self):
        records = [
            {"service_code": "NEG01", "description": "Negative Rate", "rate": -150.0},
            {"service_code": "STR02", "description": "String Rate", "rate": "InvalidPrice"},
            {"service_code": "POS03", "description": "Valid Positive Rate", "rate": 1500.0}
        ]
        validator = SOCValidator()
        report = validator.validate_batch(records)
        self.assertEqual(report["valid_count"], 1)
        self.assertEqual(report["invalid_count"], 2)

    # -------------------------------------------------------------
    # TEST 10: Invalid Dates (Start > End & Malformed)
    # -------------------------------------------------------------
    def test_10_invalid_dates(self):
        records = [
            {"service_code": "DT01", "description": "Inverted Range", "rate": 500, "effective_from": "2026-12-31", "effective_to": "2026-01-01"},
            {"service_code": "DT02", "description": "Malformed Date", "rate": 600, "effective_from": "31/12/2026"},
            {"service_code": "DT03", "description": "Valid ISO Range", "rate": 700, "effective_from": "2026-04-01", "effective_to": "2027-03-31"}
        ]
        validator = SOCValidator()
        report = validator.validate_batch(records)
        self.assertEqual(report["valid_count"], 1)
        self.assertEqual(report["invalid_count"], 2)

    # -------------------------------------------------------------
    # TEST 11: Unsupported File Format
    # -------------------------------------------------------------
    def test_11_unsupported_format(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "unsupported.txt")
        with open(file_path, "w") as f:
            f.write("Some arbitrary text")

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "error")
        self.assertIn("Unsupported file extension", res["message"])

    # -------------------------------------------------------------
    # TEST 12: Empty Document
    # -------------------------------------------------------------
    def test_12_empty_document(self):
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_12_empty.xlsx")
        wb = openpyxl.Workbook()
        wb.save(file_path)

        res = self.manager.process_file(file_path)
        self.assertEqual(res["status"], "error")
        self.assertIn("empty", res["message"].lower())

    # -------------------------------------------------------------
    # TEST 13: End-to-End Tariff Module Commit & Audit Log Verification
    # -------------------------------------------------------------
    def test_13_end_to_end_tariff_commit(self):
        # Create Excel
        file_path = os.path.join(TEST_OUTPUT_DIR, "test_13_e2e.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["CODE", "SERVICE NAME", "DEPARTMENT", "RATE"])
        ws.append(["E2E-001", "ENDOSCOPY UPPER GI", "Gastroenterology", 4500])
        ws.append(["E2E-002", "COLONOSCOPY DIAGNOSTIC", "Gastroenterology", 7500])
        wb.save(file_path)

        # Parse
        parse_res = self.manager.process_file(file_path)
        self.assertEqual(parse_res["status"], "success")
        valid_records = parse_res["validation"]["valid_records"]
        self.assertEqual(len(valid_records), 2)

        # Commit to database
        commit_res = self.manager.commit_to_tariff_module(
            valid_records=valid_records,
            file_name="test_13_e2e.xlsx",
            user="Auditor_Test",
            soc_name="TEST_E2E_SOC"
        )
        self.assertEqual(commit_res["status"], "success")
        self.assertEqual(commit_res["inserted_count"], 2)

        # Verify in SQLite tables directly
        conn = sqlite3.connect(self.test_db_path)
        cursor = conn.cursor()

        # Check tbl_soc_master
        cursor.execute("SELECT ServiceID, ServiceName, StandardRate FROM tbl_soc_master WHERE SOCName = 'TEST_E2E_SOC'")
        soc_rows = cursor.fetchall()
        self.assertEqual(len(soc_rows), 2)
        self.assertEqual(soc_rows[0][0], "E2E-001")
        self.assertEqual(soc_rows[0][2], 4500.0)

        # Check tbl_tariff_master
        cursor.execute("SELECT ServiceID, Rate FROM tbl_tariff_master WHERE ServiceID = 'E2E-002'")
        t_row = cursor.fetchone()
        self.assertIsNotNone(t_row)
        self.assertEqual(t_row[1], 7500.0)

        # Check tbl_import_logs
        cursor.execute("SELECT FileName, TotalRecords, SuccessCount, Status FROM tbl_import_logs")
        log_row = cursor.fetchone()
        self.assertIsNotNone(log_row)
        self.assertEqual(log_row[0], "test_13_e2e.xlsx")
        self.assertEqual(log_row[1], 2)
        self.assertEqual(log_row[2], 2)
        self.assertEqual(log_row[3], "Completed")

        conn.close()

if __name__ == "__main__":
    unittest.main()
